"""Render model-authored structure into safe, readable DOCX and PDF files."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import os

from app.schemas.chat_documents import DocumentSpec

MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
MAX_FILE_BYTES = 5_000_000


def render_document(spec: DocumentSpec, format: str) -> bytes:
    if format not in MIME_TYPES:
        raise ValueError("Only DOCX and PDF are supported")
    # Revalidate even when callers supplied an already constructed instance.
    spec = DocumentSpec.model_validate(spec.model_dump())
    data = _docx(spec) if format == "docx" else _pdf(spec)
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Generated file exceeds the size limit")
    return data


def _docx(spec):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (("Normal", 11), ("Title", 22), ("Heading 1", 15)):
        style = document.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    for border in list(document.styles["Title"].element.xpath("./w:pPr/w:pBdr")):
        border.getparent().remove(border)
    document.add_paragraph(spec.title, "Title")
    for block in spec.blocks:
        if block.kind == "page_break":
            document.add_page_break()
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=len(block.rows[0]))
            table.style = "Table Grid"
            for index, row in enumerate(block.rows):
                cells = table.add_row().cells
                no_split = OxmlElement("w:cantSplit")
                table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                for cell, value in zip(cells, row):
                    cell.text = value
                    if index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                if index == 0:
                    header = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(header)
        elif block.kind.endswith("list"):
            style = "List Bullet" if block.kind == "bullet_list" else "List Number"
            for value in block.items:
                document.add_paragraph(value, style)
        else:
            paragraph = document.add_paragraph(block.text, "Heading 1" if block.kind == "heading" else "Normal")
            if block.kind == "code":
                for run in paragraph.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = spec.title
    document.core_properties.author = ""
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, XPreformatted
    from reportlab.platypus.doctemplate import LayoutError

    fonts = Path(os.environ.get("CHAT_DOCUMENT_FONT_DIR", "/usr/share/fonts/truetype/dejavu"))
    for name, file in (("DocumentSans", "DejaVuSans.ttf"), ("DocumentBold", "DejaVuSans-Bold.ttf"), ("DocumentMono", "DejaVuSansMono.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / file)))
    def validate_glyphs(value, font):
        glyphs = pdfmetrics.getFont(font).face.charToGlyph
        if any(not character.isspace() and not glyphs.get(ord(character)) for character in value):
            raise ValueError("PDF font does not support this text; use DOCX to preserve it")
    validate_glyphs(spec.title, "DocumentBold")
    for block in spec.blocks:
        font = "DocumentBold" if block.kind == "heading" else "DocumentMono" if block.kind == "code" else "DocumentSans"
        for value in [block.text, *block.items, *(cell for row in block.rows for cell in row)]:
            validate_glyphs(value, font)
        if block.kind == "code" and any(pdfmetrics.stringWidth(line, font, 9) > 468 for line in block.text.expandtabs().splitlines()):
            raise ValueError("PDF code lines cannot fit the page; use DOCX or shorter lines")
    normal = ParagraphStyle("Body", fontName="DocumentSans", fontSize=10.5, leading=15, spaceAfter=8, wordWrap="CJK")
    heading = ParagraphStyle("Heading", parent=normal, fontName="DocumentBold", fontSize=15, leading=20, spaceBefore=10, keepWithNext=True)
    title = ParagraphStyle("Title", parent=heading, fontSize=22, leading=28, spaceAfter=16)
    code = ParagraphStyle("Code", parent=normal, fontName="DocumentMono", fontSize=9, leading=12)
    cell_style = ParagraphStyle("Cell", parent=normal, fontSize=9, leading=12, spaceAfter=0)
    def text(value):
        return escape(value).replace("\n", "<br/>")
    story = [Paragraph(text(spec.title), title)]
    for block in spec.blocks:
        if block.kind == "page_break":
            story.append(PageBreak())
        elif block.kind == "table":
            rows = [[Paragraph(text(cell), cell_style) for cell in row] for row in block.rows]
            table = Table(rows, colWidths=[481 / len(rows[0])] * len(rows[0]), repeatRows=1, splitInRow=0)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#bbbbbb")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([table, Spacer(1, 10)])
        elif block.kind.endswith("list"):
            for index, item in enumerate(block.items, 1):
                prefix = "• " if block.kind == "bullet_list" else f"{index}. "
                story.append(Paragraph(text(prefix + item), normal))
        elif block.kind == "code":
            story.append(XPreformatted(escape(block.text.expandtabs()), code))
        else:
            story.append(Paragraph(text(block.text), heading if block.kind == "heading" else normal))
    output = BytesIO()
    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("DocumentSans", 9)
        canvas.drawRightString(A4[0] - 57, 30, str(doc.page))
        canvas.restoreState()
    try:
        SimpleDocTemplate(output, pagesize=A4, topMargin=57, bottomMargin=57, leftMargin=57, rightMargin=57, title=spec.title, author="").build(story, onFirstPage=page_number, onLaterPages=page_number)
    except LayoutError as exc:
        raise ValueError("PDF content cannot fit a page; use DOCX or shorter table cells") from exc
    return output.getvalue()
